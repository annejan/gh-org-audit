#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Anne Jan Brouwer
# SPDX-License-Identifier: MIT

"""Dump GitHub organisation authorisations to an Excel workbook.

Sheets:
  Repos          one row per repo: archived, visibility, teams per permission
                 level, CODEOWNERS teams/users
  Matrix         repos x teams, cell = permission (A/M/W/T/R), "+CO" if the
                 team is also in CODEOWNERS
  Teams          one row per (team, member)
  Team summary   team, description, member count, repo count
  Collaborators  direct (non-team) collaborators per repo

Token: --token, or $GITHUB_TOKEN, or `gh auth token`. Needs read:org + repo.

Usage:
  python gh_org_audit.py minvws -o minvws-auth.xlsx
"""

import argparse
import base64
import os
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

API = "https://api.github.com"

# GitHub permission name -> (sort order, short code, column label)
PERM_LEVELS = {
    "admin": (0, "A", "Admin"),
    "maintain": (1, "M", "Maintain"),
    "push": (2, "W", "Write"),
    "triage": (3, "T", "Triage"),
    "pull": (4, "R", "Read"),
}

CODEOWNERS_PATHS = ["CODEOWNERS", ".github/CODEOWNERS", "docs/CODEOWNERS"]


def get_token(cli_token):
    if cli_token:
        return cli_token
    if os.environ.get("GITHUB_TOKEN"):
        return os.environ["GITHUB_TOKEN"]
    try:
        return subprocess.check_output(["gh", "auth", "token"], text=True).strip()
    except (OSError, subprocess.CalledProcessError):
        sys.exit("No token: pass --token, set GITHUB_TOKEN, or `gh auth login`")


class GitHub:
    def __init__(self, token):
        self.s = requests.Session()
        self.s.headers.update(
            {
                "Authorization": f"Bearer {token}",
                "Accept": "application/vnd.github+json",
                "X-GitHub-Api-Version": "2022-11-28",
            }
        )

    def get(self, url, params=None, ok404=False):
        if not url.startswith("http"):
            url = API + url
        for attempt in range(6):
            r = self.s.get(url, params=params, timeout=60)
            if r.status_code == 404 and ok404:
                return None
            if r.status_code in (403, 429):
                reset = r.headers.get("X-RateLimit-Reset")
                remaining = r.headers.get("X-RateLimit-Remaining")
                if remaining == "0" and reset:
                    wait = max(int(reset) - int(time.time()), 0) + 2
                    print(f"  rate limited, sleeping {wait}s", file=sys.stderr)
                    time.sleep(wait)
                    continue
                retry = r.headers.get("Retry-After")
                if retry:
                    time.sleep(int(retry))
                    continue
            if r.status_code >= 500:
                time.sleep(2**attempt)
                continue
            r.raise_for_status()
            return r
        r.raise_for_status()

    def paginate(self, url, params=None):
        params = dict(params or {}, per_page=100)
        out = []
        while url:
            r = self.get(url, params=params)
            out.extend(r.json())
            url = r.links.get("next", {}).get("url")
            params = None  # next-url already carries them
        return out


def parse_codeowners(text, org):
    """Return (teams, users) referenced in a CODEOWNERS file.

    Teams are returned as slugs (without the @org/ prefix) when they belong to
    `org`; teams of other orgs keep the full @org/slug form.
    """
    teams, users = [], []
    for line in text.splitlines():
        line = line.split("#", 1)[0].strip()
        if not line:
            continue
        for owner in line.split()[1:]:
            if owner.startswith("@"):
                if "/" in owner:
                    o, slug = owner[1:].split("/", 1)
                    key = slug if o.lower() == org.lower() else owner
                    if key not in teams:
                        teams.append(key)
                else:
                    if owner[1:] not in users:
                        users.append(owner[1:])
            elif "@" in owner:  # email address
                if owner not in users:
                    users.append(owner)
    return teams, users


def fetch_codeowners(gh, org, repo, default_branch):
    for path in CODEOWNERS_PATHS:
        r = gh.get(
            f"/repos/{org}/{repo}/contents/{path}",
            params={"ref": default_branch},
            ok404=True,
        )
        if r is None:
            continue
        data = r.json()
        if data.get("encoding") == "base64":
            text = base64.b64decode(data["content"]).decode("utf-8", "replace")
        else:
            text = gh.get(data["download_url"]).text
        return path, text
    return None, None


ROLE_TO_PERM = {"admin": "admin", "maintain": "maintain", "write": "push", "triage": "triage", "read": "pull"}


def perm_from_flags(permissions):
    """Highest permission level from a GitHub `permissions` flags object."""
    p = permissions or {}
    for level in ("admin", "maintain", "push", "triage"):
        if p.get(level):
            return level
    return "pull"


def perm_from_role(role_name, permissions):
    """Map a team repo `role_name` to a permission; custom roles fall back to flags."""
    return ROLE_TO_PERM.get((role_name or "").lower()) or perm_from_flags(permissions)


def collect_repo(gh, org, repo, team_perms):
    """team_perms: {team_slug: permission} for this repo, built from the
    team side (/orgs/{org}/teams/{slug}/repos) because /repos/{o}/{r}/teams
    needs admin rights on the repo."""
    name = repo["name"]
    perms = dict(team_perms)

    co_path, co_text = (None, None)
    try:
        co_path, co_text = fetch_codeowners(gh, org, name, repo["default_branch"])
    except requests.HTTPError as e:
        print(f"  {name}: CODEOWNERS fetch failed: {e}", file=sys.stderr)
    co_teams, co_users = parse_codeowners(co_text, org) if co_text else ([], [])

    collabs = []
    try:
        for c in gh.paginate(f"/repos/{org}/{name}/collaborators", params={"affiliation": "direct"}):
            collabs.append((c["login"], perm_from_flags(c.get("permissions"))))
    except requests.HTTPError as e:
        print(f"  {name}: collaborators fetch failed: {e}", file=sys.stderr)

    return {
        "name": name,
        "archived": repo.get("archived", False),
        "visibility": repo.get("visibility") or ("private" if repo["private"] else "public"),
        "default_branch": repo["default_branch"],
        "pushed_at": repo.get("pushed_at"),
        "team_perms": perms,
        "codeowners_path": co_path,
        "codeowners_teams": co_teams,
        "codeowners_users": co_users,
        "collaborators": collabs,
    }


def collect_team(gh, org, team):
    slug = team["slug"]
    members = gh.paginate(f"/orgs/{org}/teams/{slug}/members")
    maintainers = {m["login"] for m in gh.paginate(f"/orgs/{org}/teams/{slug}/members", params={"role": "maintainer"})}
    rows = [{"login": m["login"], "role": "maintainer" if m["login"] in maintainers else "member"} for m in members]
    repos = {}
    for r in gh.paginate(f"/orgs/{org}/teams/{slug}/repos"):
        if r["owner"]["login"].lower() != org.lower():
            continue
        repos[r["name"]] = perm_from_role(r.get("role_name"), r.get("permissions"))
    return {
        "repos": repos,
        "slug": slug,
        "name": team["name"],
        "description": team.get("description") or "",
        "privacy": team.get("privacy", ""),
        "parent": (team.get("parent") or {}).get("slug", ""),
        "members": rows,
    }


# ---------------------------------------------------------------- Excel ----

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
ARCHIVED_FILL = PatternFill("solid", fgColor="EDEDED")
PERM_FILL = {
    "A": PatternFill("solid", fgColor="F8CBAD"),
    "M": PatternFill("solid", fgColor="FFE699"),
    "W": PatternFill("solid", fgColor="C6E0B4"),
    "T": PatternFill("solid", fgColor="DDEBF7"),
    "R": PatternFill("solid", fgColor="E2EFDA"),
}


def write_sheet(ws, headers, rows, widths=None, freeze="A2"):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h)
        if w is None:
            w = min(max(len(str(h)), *(len(str(r[i - 1] or "")) for r in rows or [[""] * len(headers)])) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = w


def build_workbook(org, repos, teams, out):
    wb = Workbook()
    team_names = {t["slug"]: t["name"] for t in teams}

    def join(items):
        return ", ".join(items)

    # --- Repos
    ws = wb.active
    ws.title = "Repos"
    headers = ["Repo", "Archived", "Visibility", "Default branch", "Last push"]
    headers += [f"Teams: {label}" for _, _, label in sorted(PERM_LEVELS.values())]
    headers += ["CODEOWNERS file", "CODEOWNERS teams", "CODEOWNERS users", "Direct collaborators", "Team count"]
    rows = []
    for r in repos:
        by_level = {lvl: [] for lvl in PERM_LEVELS}
        for slug, p in sorted(r["team_perms"].items()):
            by_level.setdefault(p, []).append(slug)
        row = [
            r["name"],
            "yes" if r["archived"] else "no",
            r["visibility"],
            r["default_branch"],
            (r["pushed_at"] or "")[:10],
        ]
        for lvl, _ in sorted(PERM_LEVELS.items(), key=lambda kv: kv[1][0]):
            row.append(join(by_level.get(lvl, [])))
        row += [
            r["codeowners_path"] or "",
            join(r["codeowners_teams"]),
            join(r["codeowners_users"]),
            join(f"{login} ({PERM_LEVELS[lvl][1]})" for login, lvl in r["collaborators"]),
            len(r["team_perms"]),
        ]
        rows.append(row)
    write_sheet(ws, headers, rows, widths={"Repo": 40})
    for i, r in enumerate(repos, 2):
        if r["archived"]:
            for c in ws[i]:
                c.fill = ARCHIVED_FILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    # --- Matrix
    ws = wb.create_sheet("Matrix")
    slugs = sorted(team_names, key=str.lower)
    # include teams that appear on repos but not in org listing (shouldn't happen)
    for r in repos:
        for s in list(r["team_perms"]) + r["codeowners_teams"]:
            if s not in team_names and s not in slugs:
                slugs.append(s)
    headers = ["Repo", "Archived", "Visibility"] + slugs
    rows = []
    for r in repos:
        row = [r["name"], "yes" if r["archived"] else "no", r["visibility"]]
        for s in slugs:
            p = r["team_perms"].get(s)
            cell = PERM_LEVELS[p][1] if p else ""
            if s in r["codeowners_teams"]:
                cell = (cell + "+CO") if cell else "CO"
            row.append(cell)
        rows.append(row)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(text_rotation=90, vertical="bottom")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 9
    ws.column_dimensions["C"].width = 10
    for i in range(4, 4 + len(slugs)):
        ws.column_dimensions[get_column_letter(i)].width = 4.5
    ws.row_dimensions[1].height = 140
    for row in ws.iter_rows(min_row=2, min_col=4):
        for c in row:
            if c.value:
                c.alignment = Alignment(horizontal="center")
                c.fill = PERM_FILL.get(str(c.value)[0], PatternFill())
    ws.cell(
        row=len(rows) + 3, column=1, value="A=Admin  M=Maintain  W=Write  T=Triage  R=Read  CO=in CODEOWNERS"
    ).font = Font(italic=True)

    # --- Teams (members)
    ws = wb.create_sheet("Teams")
    headers = ["Team", "Team name", "Member", "Role in team", "Parent team"]
    rows = []
    for t in sorted(teams, key=lambda t: t["slug"].lower()):
        if not t["members"]:
            rows.append([t["slug"], t["name"], "(no members)", "", t["parent"]])
        for m in sorted(t["members"], key=lambda m: m["login"].lower()):
            rows.append([t["slug"], t["name"], m["login"], m["role"], t["parent"]])
    write_sheet(ws, headers, rows)

    # --- Team summary
    ws = wb.create_sheet("Team summary")
    headers = [
        "Team",
        "Team name",
        "Description",
        "Privacy",
        "Parent team",
        "Members",
        "Repos (any access)",
        "Repos admin",
        "Repos write+",
        "Repos read-only",
        "CODEOWNERS in repos",
    ]
    rows = []
    for t in sorted(teams, key=lambda t: t["slug"].lower()):
        s = t["slug"]
        with_access = [r for r in repos if s in r["team_perms"]]
        rows.append(
            [
                s,
                t["name"],
                t["description"],
                t["privacy"],
                t["parent"],
                len(t["members"]),
                len(with_access),
                sum(1 for r in with_access if r["team_perms"][s] == "admin"),
                sum(1 for r in with_access if r["team_perms"][s] in ("maintain", "push")),
                sum(1 for r in with_access if r["team_perms"][s] in ("triage", "pull")),
                sum(1 for r in repos if s in r["codeowners_teams"]),
            ]
        )
    write_sheet(ws, headers, rows, widths={"Description": 50})

    # --- Members -> teams (handy reverse lookup)
    ws = wb.create_sheet("Members")
    by_login = {}
    for t in teams:
        for m in t["members"]:
            by_login.setdefault(m["login"], []).append(t["slug"])
    headers = ["Member", "Team count", "Teams"]
    rows = [[login, len(ts), join(sorted(ts))] for login, ts in sorted(by_login.items(), key=lambda kv: kv[0].lower())]
    write_sheet(ws, headers, rows, widths={"Teams": 80})

    # --- Collaborators
    ws = wb.create_sheet("Collaborators")
    headers = ["Repo", "Archived", "Collaborator", "Permission"]
    rows = []
    for r in repos:
        for login, lvl in r["collaborators"]:
            rows.append([r["name"], "yes" if r["archived"] else "no", login, PERM_LEVELS[lvl][2]])
    write_sheet(ws, headers, rows)

    # --- Info
    ws = wb.create_sheet("Info")
    ws.append(["Organisation", org])
    ws.append(["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    ws.append(["Repos", len(repos)])
    ws.append(["Teams", len(teams)])
    ws.append([])
    ws.append(["Legend", "A=Admin  M=Maintain  W=Write (push)  T=Triage  R=Read (pull)"])
    ws.append(
        [
            "",
            "CODEOWNERS teams: teams referenced as @org/team in CODEOWNERS (any of "
            + ", ".join(CODEOWNERS_PATHS)
            + " on default branch)",
        ]
    )
    ws.append(
        [
            "",
            "Team permissions come from GitHub API /orgs/{org}/teams/{team}/repos (role_name). "
            "Secret teams you are not a member of are invisible and therefore missing.",
        ]
    )
    ws.append(["", "Direct collaborators need admin rights on the repo to be listed; otherwise the column is empty."])
    ws.append(["", "Direct collaborators = users granted access on the repo directly, outside a team."])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 110

    wb.save(out)


# ----------------------------------------------------------------- main ----


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("org")
    ap.add_argument("-o", "--output", help="xlsx path (default <org>-authorisations-<date>.xlsx)")
    ap.add_argument("--token")
    ap.add_argument("-j", "--jobs", type=int, default=8, help="parallel API workers (default 8)")
    ap.add_argument("--skip-archived", action="store_true", help="leave archived repos out")
    args = ap.parse_args()

    out = args.output or f"{args.org}-authorisations-{datetime.now():%Y-%m-%d}.xlsx"
    gh = GitHub(get_token(args.token))

    print(f"Listing repos of {args.org} ...", file=sys.stderr)
    repos_raw = gh.paginate(f"/orgs/{args.org}/repos", params={"type": "all"})
    if args.skip_archived:
        repos_raw = [r for r in repos_raw if not r.get("archived")]
    repos_raw.sort(key=lambda r: r["name"].lower())
    print(f"  {len(repos_raw)} repos", file=sys.stderr)

    print("Listing teams ...", file=sys.stderr)
    teams_raw = gh.paginate(f"/orgs/{args.org}/teams")
    print(f"  {len(teams_raw)} teams", file=sys.stderr)

    with ThreadPoolExecutor(max_workers=args.jobs) as ex:
        print("Fetching team members ...", file=sys.stderr)
        teams = list(ex.map(lambda t: collect_team(gh, args.org, t), teams_raw))

        perms_by_repo = {}
        for t in teams:
            for rname, perm in t["repos"].items():
                perms_by_repo.setdefault(rname, {})[t["slug"]] = perm

        print("Fetching CODEOWNERS / collaborators ...", file=sys.stderr)
        repos = []
        for i, res in enumerate(
            ex.map(lambda r: collect_repo(gh, args.org, r, perms_by_repo.get(r["name"], {})), repos_raw), 1
        ):
            repos.append(res)
            if i % 25 == 0 or i == len(repos_raw):
                print(f"  {i}/{len(repos_raw)}", file=sys.stderr)

    build_workbook(args.org, repos, teams, out)
    rl = gh.get("/rate_limit").json()["resources"]["core"]
    print(f"Wrote {out}  (rate limit left: {rl['remaining']}/{rl['limit']})", file=sys.stderr)


if __name__ == "__main__":
    main()
