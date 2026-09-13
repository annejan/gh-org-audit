# SPDX-FileCopyrightText: 2026 Anne Jan Brouwer
# SPDX-License-Identifier: MIT

import subprocess
import sys

import pytest
from openpyxl import load_workbook

import gh_org_audit
from gh_org_audit import get_token
from tests.conftest import FakeResponse, FakeSession


def test_get_token_precedence(monkeypatch):
    monkeypatch.setenv("GITHUB_TOKEN", "env-tok")
    assert get_token("cli-tok") == "cli-tok"
    assert get_token(None) == "env-tok"
    monkeypatch.delenv("GITHUB_TOKEN")
    monkeypatch.setattr(subprocess, "check_output", lambda *a, **k: "gh-tok\n")
    assert get_token(None) == "gh-tok"


def test_get_token_exits_without_gh(monkeypatch):
    monkeypatch.delenv("GITHUB_TOKEN", raising=False)

    def boom(*a, **k):
        raise FileNotFoundError

    monkeypatch.setattr(subprocess, "check_output", boom)
    with pytest.raises(SystemExit):
        get_token(None)


def test_main_end_to_end(tmp_path, monkeypatch):
    session = FakeSession()
    session.add(
        "/orgs/o/repos",
        FakeResponse(
            200,
            [
                {"name": "z", "private": False, "default_branch": "main", "archived": True},
                {"name": "a", "private": True, "default_branch": "main", "archived": False},
            ],
        ),
        params={"type": "all"},
    )
    session.add("/orgs/o/teams", FakeResponse(200, [{"slug": "t", "name": "T"}]))
    session.add("/orgs/o/teams/t/members", FakeResponse(200, [{"login": "alice"}]))
    session.add("/orgs/o/teams/t/members", FakeResponse(200, []), params={"role": "maintainer"})
    session.add(
        "/orgs/o/teams/t/repos",
        FakeResponse(
            200,
            [
                {"name": "a", "owner": {"login": "o"}, "role_name": "admin", "permissions": {}},
            ],
        ),
    )
    session.add("/rate_limit", FakeResponse(200, {"resources": {"core": {"remaining": 1, "limit": 2}}}))

    monkeypatch.setattr(gh_org_audit.GitHub, "__init__", lambda self, token: setattr(self, "s", session))
    out = tmp_path / "o.xlsx"
    monkeypatch.setattr(sys, "argv", ["gh-org-audit", "o", "-o", str(out), "--token", "x", "-j", "1"])
    gh_org_audit.main()

    wb = load_workbook(out)
    names = [r[0] for r in wb["Repos"].iter_rows(min_row=2, values_only=True)]
    assert names == ["a", "z"]  # sorted case-insensitively
    assert wb["Matrix"]["D2"].value == "A"


def test_main_skip_archived(tmp_path, monkeypatch):
    session = FakeSession()
    session.add(
        "/orgs/o/repos",
        FakeResponse(
            200,
            [
                {"name": "z", "private": False, "default_branch": "main", "archived": True},
                {"name": "a", "private": True, "default_branch": "main", "archived": False},
            ],
        ),
        params={"type": "all"},
    )
    session.add("/orgs/o/teams", FakeResponse(200, []))
    session.add("/rate_limit", FakeResponse(200, {"resources": {"core": {"remaining": 1, "limit": 2}}}))
    monkeypatch.setattr(gh_org_audit.GitHub, "__init__", lambda self, token: setattr(self, "s", session))
    out = tmp_path / "o.xlsx"
    monkeypatch.setattr(sys, "argv", ["gh-org-audit", "o", "-o", str(out), "--token", "x", "--skip-archived"])
    gh_org_audit.main()
    names = [r[0] for r in load_workbook(out)["Repos"].iter_rows(min_row=2, values_only=True)]
    assert names == ["a"]
