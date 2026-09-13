# SPDX-FileCopyrightText: 2026 Anne Jan Brouwer
# SPDX-License-Identifier: MIT

import pytest
from openpyxl import load_workbook

from gh_org_audit import build_workbook

REPOS = [
    {
        "name": "alpha",
        "archived": False,
        "visibility": "public",
        "default_branch": "main",
        "pushed_at": "2026-05-06T07:08:09Z",
        "team_perms": {"admins": "admin", "devs": "push", "readers": "pull"},
        "codeowners_path": ".github/CODEOWNERS",
        "codeowners_teams": ["devs", "ghost-team"],
        "codeowners_users": ["carol"],
        "collaborators": [("dave", "admin")],
    },
    {
        "name": "beta",
        "archived": True,
        "visibility": "private",
        "default_branch": "master",
        "pushed_at": None,
        "team_perms": {"devs": "maintain"},
        "codeowners_path": None,
        "codeowners_teams": [],
        "codeowners_users": [],
        "collaborators": [],
    },
]

TEAMS = [
    {
        "slug": "admins",
        "name": "Admins",
        "description": "the bosses",
        "privacy": "closed",
        "parent": "",
        "members": [{"login": "alice", "role": "maintainer"}],
        "repos": {"alpha": "admin"},
    },
    {
        "slug": "devs",
        "name": "Devs",
        "description": "",
        "privacy": "closed",
        "parent": "admins",
        "members": [{"login": "alice", "role": "member"}, {"login": "bob", "role": "member"}],
        "repos": {"alpha": "push", "beta": "maintain"},
    },
    {
        "slug": "readers",
        "name": "Readers",
        "description": "",
        "privacy": "secret",
        "parent": "",
        "members": [],
        "repos": {"alpha": "pull"},
    },
]


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    out = tmp_path_factory.mktemp("xlsx") / "out.xlsx"
    build_workbook("o", REPOS, TEAMS, out)
    return load_workbook(out)


def rows(ws):
    # openpyxl reads empty strings back as None; normalise so fixtures stay readable
    return [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]


def test_sheet_names(wb):
    assert wb.sheetnames == ["Repos", "Matrix", "Teams", "Team summary", "Members", "Collaborators", "Info"]


def test_repos_sheet(wb):
    r = rows(wb["Repos"])
    hdr = r[0]
    assert hdr[:5] == ["Repo", "Archived", "Visibility", "Default branch", "Last push"]
    assert "Teams: Admin" in hdr and "Teams: Read" in hdr
    alpha = dict(zip(hdr, r[1], strict=True))
    assert alpha["Repo"] == "alpha"
    assert alpha["Archived"] == "no"
    assert alpha["Last push"] == "2026-05-06"
    assert alpha["Teams: Admin"] == "admins"
    assert alpha["Teams: Write"] == "devs"
    assert alpha["Teams: Read"] == "readers"
    assert alpha["CODEOWNERS file"] == ".github/CODEOWNERS"
    assert alpha["CODEOWNERS teams"] == "devs, ghost-team"
    assert alpha["CODEOWNERS users"] == "carol"
    assert alpha["Direct collaborators"] == "dave (A)"
    assert alpha["Team count"] == 3
    beta = dict(zip(hdr, r[2], strict=True))
    assert beta["Archived"] == "yes"
    assert beta["Last push"] == ""
    assert beta["Teams: Maintain"] == "devs"
    assert beta["Team count"] == 1


def test_repos_sheet_filter_and_freeze(wb):
    ws = wb["Repos"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None


def test_matrix_sheet(wb):
    r = rows(wb["Matrix"])
    hdr = r[0]
    # org teams sorted, plus team only seen in CODEOWNERS appended
    assert hdr == ["Repo", "Archived", "Visibility", "admins", "devs", "readers", "ghost-team"]
    assert r[1] == ["alpha", "no", "public", "A", "W+CO", "R", "CO"]
    assert r[2] == ["beta", "yes", "private", "", "M", "", ""]
    assert wb["Matrix"].freeze_panes == "D2"
    legend = [c for row in r for c in row if isinstance(c, str) and c.startswith("A=Admin")]
    assert legend


def test_teams_sheet(wb):
    r = rows(wb["Teams"])
    assert r[0] == ["Team", "Team name", "Member", "Role in team", "Parent team"]
    assert r[1] == ["admins", "Admins", "alice", "maintainer", ""]
    assert r[2] == ["devs", "Devs", "alice", "member", "admins"]
    assert r[3] == ["devs", "Devs", "bob", "member", "admins"]
    assert r[4] == ["readers", "Readers", "(no members)", "", ""]


def test_team_summary(wb):
    r = rows(wb["Team summary"])
    hdr = r[0]
    devs = dict(zip(hdr, r[2], strict=True))
    assert devs["Team"] == "devs"
    assert devs["Members"] == 2
    assert devs["Repos (any access)"] == 2
    assert devs["Repos admin"] == 0
    assert devs["Repos write+"] == 2
    assert devs["Repos read-only"] == 0
    assert devs["CODEOWNERS in repos"] == 1
    admins = dict(zip(hdr, r[1], strict=True))
    assert admins["Repos admin"] == 1
    assert admins["Description"] == "the bosses"


def test_members_sheet(wb):
    r = rows(wb["Members"])
    assert r[0] == ["Member", "Team count", "Teams"]
    assert r[1] == ["alice", 2, "admins, devs"]
    assert r[2] == ["bob", 1, "devs"]


def test_collaborators_sheet(wb):
    r = rows(wb["Collaborators"])
    assert r[0] == ["Repo", "Archived", "Collaborator", "Permission"]
    assert r[1] == ["alpha", "no", "dave", "Admin"]
    assert len(r) == 2


def test_info_sheet(wb):
    r = rows(wb["Info"])
    assert r[0] == ["Organisation", "o"]
    assert r[2] == ["Repos", 2]
    assert r[3] == ["Teams", 3]


def test_empty_org(tmp_path):
    out = tmp_path / "empty.xlsx"
    build_workbook("o", [], [], out)
    wb = load_workbook(out)
    assert rows(wb["Repos"])[0][0] == "Repo"
    assert len(rows(wb["Teams"])) == 1
